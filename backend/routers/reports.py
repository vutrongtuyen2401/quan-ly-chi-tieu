"""Báo cáo tổng quan, xuất CSV/Excel, xu hướng, theo tuần, so sánh tháng."""

import csv
import datetime
import io
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from backend.db import get_db
from backend.schemas import DATE_PATTERN, MONTH_PATTERN
from backend.security import get_current_user
from backend.services import process_recurring_transactions

router = APIRouter()


# ──────────────────────────────────────────────
# REPORTS ROUTES
# ──────────────────────────────────────────────
@router.get("/api/reports/summary")
def get_reports_summary(month_year: Optional[str] = Query(None, pattern=MONTH_PATTERN), user: dict = Depends(get_current_user)):
    if not month_year:
        month_year = datetime.date.today().strftime("%Y-%m")
    with get_db() as conn:
        process_recurring_transactions(conn, user["user_id"])
        income = conn.execute("""
            SELECT COALESCE(SUM(amount), 0) as total FROM transactions
            WHERE user_id = ? AND transaction_type = 'INCOME'
            AND strftime('%Y-%m', transaction_date) = ?
        """, (user["user_id"], month_year)).fetchone()["total"]

        expense = conn.execute("""
            SELECT COALESCE(SUM(amount), 0) as total FROM transactions
            WHERE user_id = ? AND transaction_type = 'EXPENSE'
            AND strftime('%Y-%m', transaction_date) = ?
        """, (user["user_id"], month_year)).fetchone()["total"]

        total_balance = conn.execute("""
            SELECT COALESCE(SUM(balance), 0) as total FROM wallets WHERE user_id = ?
        """, (user["user_id"],)).fetchone()["total"]

        # Chi tiêu theo danh mục
        by_category = conn.execute("""
            SELECT c.category_name, c.icon, SUM(t.amount) as total
            FROM transactions t
            LEFT JOIN categories c ON t.category_id = c.id
            WHERE t.user_id = ? AND t.transaction_type = 'EXPENSE'
            AND strftime('%Y-%m', t.transaction_date) = ?
            GROUP BY t.category_id
            ORDER BY total DESC
        """, (user["user_id"], month_year)).fetchall()

        return {
            "month_year": month_year,
            "total_income": income,
            "total_expense": expense,
            "net_savings": income - expense,
            "total_balance": total_balance,
            "expense_by_category": [dict(r) for r in by_category],
        }


# ──────────────────────────────────────────────
# Change 7: EXPORT REPORTS (CSV / EXCEL)
# ──────────────────────────────────────────────
@router.get("/api/reports/export")
def export_reports(
    start_date: Optional[str] = Query(None, pattern=DATE_PATTERN),
    end_date: Optional[str] = Query(None, pattern=DATE_PATTERN),
    format: str = Query("csv", pattern="^(csv|excel)$"),
    user: dict = Depends(get_current_user)
):
    """Xuất lịch sử thu chi ra file CSV hoặc Excel"""
    with get_db() as conn:
        where_clauses = ["t.user_id = ?"]
        params = [user["user_id"]]

        if start_date:
            where_clauses.append("t.transaction_date >= ?")
            params.append(start_date)
        if end_date:
            where_clauses.append("t.transaction_date <= ?")
            params.append(end_date)

        where_sql = " AND ".join(where_clauses)
        rows = conn.execute(f"""
            SELECT t.transaction_date, c.category_name, t.transaction_type,
                   t.amount, w.wallet_name, t.note
            FROM transactions t
            LEFT JOIN wallets w ON t.wallet_id = w.id
            LEFT JOIN categories c ON t.category_id = c.id
            WHERE {where_sql}
            ORDER BY t.transaction_date DESC, t.id DESC
        """, params).fetchall()

    today_str = datetime.date.today().strftime("%Y%m%d")

    if format == "csv":
        output = io.StringIO()
        output.write('\ufeff')  # UTF-8 BOM
        writer = csv.writer(output)
        writer.writerow(["Ngày Giao Dịch", "Danh Mục", "Loại Giao Dịch", "Số Tiền (VNĐ)", "Túi / Ví", "Ghi Chú"])

        for r in rows:
            t_type_str = "Thu Nhập" if r["transaction_type"] == "INCOME" else "Chi Tiêu"
            writer.writerow([
                r["transaction_date"],
                r["category_name"] or "Không rõ",
                t_type_str,
                f"{r['amount']:,.0f}",
                r["wallet_name"] or "Không rõ",
                r["note"] or ""
            ])

        csv_bytes = output.getvalue().encode("utf-8-sig")
        return StreamingResponse(
            io.BytesIO(csv_bytes),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=bao_cao_chi_tieu_{today_str}.csv"}
        )

    else:  # format == "excel"
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Báo Cáo Chi Tiêu"

        headers = ["Ngày Giao Dịch", "Danh Mục", "Loại Giao Dịch", "Số Tiền (VNĐ)", "Túi / Ví", "Ghi Chú"]
        ws.append(headers)

        header_fill = PatternFill(start_color="2B8A82", end_color="2B8A82", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in rows:
            t_type_str = "Thu Nhập" if r["transaction_type"] == "INCOME" else "Chi Tiêu"
            ws.append([
                r["transaction_date"],
                r["category_name"] or "Không rõ",
                t_type_str,
                r["amount"],
                r["wallet_name"] or "Không rõ",
                r["note"] or ""
            ])

        for row_idx in range(2, len(rows) + 2):
            amount_cell = ws.cell(row=row_idx, column=4)
            amount_cell.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        return StreamingResponse(
            excel_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=bao_cao_chi_tieu_{today_str}.xlsx"}
        )


# ──────────────────────────────────────────────
# ADVANCED REPORTS
# ──────────────────────────────────────────────
@router.get("/api/reports/trend")
def get_trend_report(months: int = Query(6, ge=1, le=12), user: dict = Depends(get_current_user)):
    """Xu hướng thu/chi N tháng gần nhất (tính cả tháng hiện tại; tháng không có giao dịch vẫn trả về 0)"""
    today = datetime.date.today()
    year, month = today.year, today.month - (months - 1)
    while month <= 0:
        year, month = year - 1, month + 12
    start = datetime.date(year, month, 1)
    next_month_start = (datetime.date(today.year + 1, 1, 1) if today.month == 12
                        else datetime.date(today.year, today.month + 1, 1))

    with get_db() as conn:
        rows = conn.execute("""
            SELECT strftime('%Y-%m', transaction_date) as month,
                   COALESCE(SUM(CASE WHEN transaction_type='INCOME' THEN amount ELSE 0 END), 0) as income,
                   COALESCE(SUM(CASE WHEN transaction_type='EXPENSE' THEN amount ELSE 0 END), 0) as expense
            FROM transactions
            WHERE user_id = ? AND transaction_date >= ? AND transaction_date < ?
            GROUP BY month
        """, (user["user_id"], start.isoformat(), next_month_start.isoformat())).fetchall()

    by_month = {r["month"]: r for r in rows}
    result = []
    y, m = start.year, start.month
    for _ in range(months):
        key = f"{y}-{m:02d}"
        r = by_month.get(key)
        income = r["income"] if r else 0
        expense = r["expense"] if r else 0
        result.append({"month": key, "income": income, "expense": expense, "savings": income - expense})
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)

    return {"months": months, "trend": result}


@router.get("/api/reports/weekly")
def get_weekly_report(weeks: int = Query(4, ge=1, le=12), user: dict = Depends(get_current_user)):
    """Thu/chi theo tuần (tuần bắt đầu từ thứ Hai, N tuần gần nhất tính cả tuần hiện tại)"""
    today = datetime.date.today()
    this_monday = today - datetime.timedelta(days=today.weekday())
    start = this_monday - datetime.timedelta(weeks=weeks - 1)
    end = this_monday + datetime.timedelta(days=7)

    with get_db() as conn:
        rows = conn.execute("""
            SELECT date(transaction_date,
                        printf('-%d days', (CAST(strftime('%w', transaction_date) AS INTEGER) + 6) % 7)) as week_start,
                   COALESCE(SUM(CASE WHEN transaction_type='EXPENSE' THEN amount ELSE 0 END), 0) as expense,
                   COALESCE(SUM(CASE WHEN transaction_type='INCOME' THEN amount ELSE 0 END), 0) as income,
                   COUNT(*) as txn_count
            FROM transactions
            WHERE user_id = ? AND transaction_date >= ? AND transaction_date < ?
            GROUP BY week_start
        """, (user["user_id"], start.isoformat(), end.isoformat())).fetchall()

    by_week = {r["week_start"]: r for r in rows}
    data = []
    for i in range(weeks):
        monday = start + datetime.timedelta(weeks=i)
        r = by_week.get(monday.isoformat())
        iso_year, iso_week, _ = monday.isocalendar()
        data.append({
            "week": f"{iso_year}-W{iso_week:02d}",
            "week_start": monday.isoformat(),
            "expense": r["expense"] if r else 0,
            "income": r["income"] if r else 0,
            "txn_count": r["txn_count"] if r else 0,
        })

    return {"weeks": weeks, "data": data}


@router.get("/api/reports/compare")
def compare_months(
    month1: str = Query(..., description="YYYY-MM", pattern=MONTH_PATTERN),
    month2: str = Query(..., description="YYYY-MM", pattern=MONTH_PATTERN),
    user: dict = Depends(get_current_user)
):
    """So sánh chi tiêu 2 tháng"""
    with get_db() as conn:
        results = {}
        for label, month in [("month1", month1), ("month2", month2)]:
            income = conn.execute("""
                SELECT COALESCE(SUM(amount), 0) as total FROM transactions
                WHERE user_id = ? AND transaction_type = 'INCOME'
                AND strftime('%Y-%m', transaction_date) = ?
            """, (user["user_id"], month)).fetchone()["total"]

            expense = conn.execute("""
                SELECT COALESCE(SUM(amount), 0) as total FROM transactions
                WHERE user_id = ? AND transaction_type = 'EXPENSE'
                AND strftime('%Y-%m', transaction_date) = ?
            """, (user["user_id"], month)).fetchone()["total"]

            by_category = conn.execute("""
                SELECT c.category_name, c.icon, SUM(t.amount) as total
                FROM transactions t
                LEFT JOIN categories c ON t.category_id = c.id
                WHERE t.user_id = ? AND t.transaction_type = 'EXPENSE'
                AND strftime('%Y-%m', t.transaction_date) = ?
                GROUP BY t.category_id ORDER BY total DESC
            """, (user["user_id"], month)).fetchall()

            results[label] = {
                "month": month,
                "income": income,
                "expense": expense,
                "savings": income - expense,
                "by_category": [dict(r) for r in by_category],
            }

        # Tính delta
        delta_income = results["month2"]["income"] - results["month1"]["income"]
        delta_expense = results["month2"]["expense"] - results["month1"]["expense"]

        return {
            **results,
            "delta_income": delta_income,
            "delta_expense": delta_expense,
            "delta_savings": (results["month2"]["savings"]) - (results["month1"]["savings"]),
        }
